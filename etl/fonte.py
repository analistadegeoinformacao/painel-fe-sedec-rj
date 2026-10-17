"""Leitura da planilha. Devolve sempre {nome_da_aba: [[célula, ...], ...]}.

Três modos:
  google  — produção (GitHub Actions). Usa uma conta de serviço com acesso
            SOMENTE LEITURA à planilha. Credencial no segredo GOOGLE_CREDENTIALS.
  xlsx    — arquivo baixado (Arquivo > Fazer download > .xlsx).
  texto   — dump em texto com blocos '## Sheet name: <aba>' (formato do
            arquivo NOVA_BASE_ANALISTAS usado para desenvolver este projeto).
"""
from __future__ import annotations

import csv
import io
import json
import os
import re
from pathlib import Path


def ler_google(planilha_id: str, abas: list[str]) -> dict[str, list[list[str]]]:
    import gspread  # importado só aqui: testes locais não precisam da lib
    from google.oauth2.service_account import Credentials

    info = json.loads(os.environ["GOOGLE_CREDENTIALS"])
    cred = Credentials.from_service_account_info(
        info, scopes=["https://www.googleapis.com/auth/spreadsheets.readonly"]
    )
    livro = gspread.authorize(cred).open_by_key(planilha_id)
    existentes = {ws.title: ws for ws in livro.worksheets()}
    saida = {}
    for aba in abas:
        if aba in existentes:
            # FORMATTED_VALUE = o mesmo texto que o analista vê na tela
            saida[aba] = existentes[aba].get_all_values(value_render_option="FORMATTED_VALUE")
    return saida


def ler_xlsx(caminho: Path) -> dict[str, list[list[str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(caminho, read_only=True, data_only=True)
    saida = {}
    for ws in wb.worksheets:
        saida[ws.title] = [["" if c is None else c for c in linha] for linha in ws.iter_rows(values_only=True)]
    return saida


def ler_texto(caminho: Path) -> dict[str, list[list[str]]]:
    txt = Path(caminho).read_text(encoding="utf-8")
    partes = re.split(r"^## Sheet name: (.*)$", txt, flags=re.M)
    saida = {}
    for i in range(1, len(partes), 2):
        corpo = partes[i + 1].lstrip("\n")
        saida[partes[i].strip()] = list(csv.reader(io.StringIO(corpo)))
    return saida


def ler(modo: str, origem: str, abas: list[str]):
    if modo == "google":
        return ler_google(origem, abas)
    if modo == "xlsx":
        return ler_xlsx(Path(origem))
    if modo == "texto":
        return ler_texto(Path(origem))
    raise SystemExit(f"modo de fonte desconhecido: {modo}")
